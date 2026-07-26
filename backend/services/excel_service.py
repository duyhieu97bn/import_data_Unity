import re
from datetime import datetime
from openpyxl import load_workbook
import config

_IFERROR_PATTERN = re.compile(r'IFERROR\(.*?,\s*(.*?)\)\s*$')

RECORD_COLS = [
    "stt", "clb", "ngay", "thang", "nam", "ngayThangNam",
    "quyenSo", "soBN", "maKH", "hoTenKH", "ngaySinh", "sdt",
    "maGoi", "sale", "ptChamSoc", "phanBo1", "loaiThe1",
    "loaiThe2", "phanBo2", "csNhap", "tongSS", "ptGym", "pttl",
    "daily", "boi", "ptTang", "buoiPtCoDinh", "soThang",
    "tongSoThang", "congThangBL", "congThangTang", "congNgayTang",
    "goiTap", "khungGioTap", "chuyenNhuong", "soHD", "soThangBuoi",
    "ngayBatDau", "ngayKetThuc", "tinhTrang", "tongSoNgay",
    "tongNgayDaDi", "tongSoNgayConLai", "quaTang", "donGiaNiemYet",
    "tongTienBaoLuu", "tongTienBuoiDay", "tongTien", "mucUuDai",
    "soTienKHGiam", "tong", "soTienDVPhaiThu", "bnDatCoc",
    "tenGoiNangCap", "soTienKHNangCap", "khDaDatCoc",
    "tongDatCocNangCap", "conPhaiThuKH", "phanTramSaleDT",
    "datCoc", "ttPhanThieu", "dongFull", "phiPos", "doanhThu",
    "hinhThucTT", "daNhanTM", "ckChung", "tongTienCKChung",
    "trangThai", "tongGoiTap", "loai", "luot", "nguon",
    "daNhapTimeSoft", "ghiChu", "daKichHoat",
    "tongHopDayShowGX", "thangHoanThienHD",
]

GOI_COLS = [
    "maGoi", "clb", "loaiThe", "maGoiHoatDong", "phanBo",
    "maLoaiTheTarget", "maLoaiTheLocPT", "thoiGian", "thoiGian2",
    "giaNiemYet", "giaNiemYetGD", "ghiChu",
]


def row_to_dict(row, cols):
    return dict(zip(cols, row)) if row else None


def safe_str(v):
    if v is None:
        return ""
    return str(v).strip()


def safe_num(v):
    if v is None or v == "" or v == "-":
        return 0
    try:
        return float(v)
    except (ValueError, TypeError):
        return 0


def parse_date(v):
    if isinstance(v, datetime):
        return v.strftime("%d/%m/%Y")
    s = safe_str(v)
    return s if s else ""


def _formula_fallback(cell):
    if cell.value is None:
        return ''
    s = str(cell.value)
    m = _IFERROR_PATTERN.search(s)
    if m:
        return m.group(1).strip().strip('"')
    return s


def _parse_goi_num(s):
    if not s or s == '-':
        return 0
    try:
        return float(str(s).replace(',', ''))
    except:
        return 0


def parse_goi_from_file(filepath: str):
    wb = load_workbook(filepath)
    ws = wb[wb.sheetnames[0]]
    items = []
    for row in range(4, ws.max_row + 1):
        ma_goi = _formula_fallback(ws.cell(row=row, column=5))
        if not ma_goi:
            continue
        items.append({
            "maGoi": ma_goi,
            "clb": _formula_fallback(ws.cell(row=row, column=2)),
            "loaiThe": _formula_fallback(ws.cell(row=row, column=3)),
            "maGoiHoatDong": _formula_fallback(ws.cell(row=row, column=4)),
            "phanBo": _formula_fallback(ws.cell(row=row, column=6)),
            "maLoaiTheTarget": _formula_fallback(ws.cell(row=row, column=7)),
            "maLoaiTheLocPT": _formula_fallback(ws.cell(row=row, column=8)),
            "thoiGian": _formula_fallback(ws.cell(row=row, column=9)),
            "thoiGian2": _formula_fallback(ws.cell(row=row, column=10)),
            "giaNiemYet": _parse_goi_num(_formula_fallback(ws.cell(row=row, column=11))),
            "giaNiemYetGD": _parse_goi_num(_formula_fallback(ws.cell(row=row, column=12))),
            "ghiChu": _formula_fallback(ws.cell(row=row, column=13)),
        })
    return items


def parse_records_from_file(filepath: str):
    wb = load_workbook(filepath, data_only=True)
    sheet_names = wb.sheetnames
    ws = None
    if config.SHEET1_NAME in sheet_names:
        ws = wb[config.SHEET1_NAME]
    else:
        ws = wb[sheet_names[0]]
    data_start = None
    for row in range(1, min(ws.max_row + 1, 20)):
        vals = [ws.cell(row=row, column=c).value for c in range(1, 20)]
        str_vals = [str(v).strip().lower() if v else "" for v in vals]
        if any("tên" in s or "họ" in s or "khách" in s or "hoten" in s for s in str_vals):
            data_start = row + 1
            break
    if data_start is None:
        data_start = 4
    records = []
    for row in range(data_start, ws.max_row + 1):
        stt = ws.cell(row=row, column=1).value
        if stt is None or str(stt).strip() == "":
            continue
        records.append({
            "stt": stt,
            "clb": safe_str(ws.cell(row=row, column=2).value),
            "ngay": ws.cell(row=row, column=3).value or "",
            "thang": ws.cell(row=row, column=4).value or "",
            "nam": ws.cell(row=row, column=5).value or "",
            "ngayThangNam": safe_str(ws.cell(row=row, column=6).value),
            "quyenSo": ws.cell(row=row, column=7).value or "",
            "soBN": ws.cell(row=row, column=8).value or "",
            "maKH": safe_str(ws.cell(row=row, column=9).value),
            "hoTenKH": safe_str(ws.cell(row=row, column=10).value),
            "ngaySinh": safe_str(ws.cell(row=row, column=11).value),
            "sdt": safe_str(ws.cell(row=row, column=12).value),
            "maGoi": safe_str(ws.cell(row=row, column=13).value),
            "sale": safe_str(ws.cell(row=row, column=14).value),
            "ptChamSoc": safe_str(ws.cell(row=row, column=15).value),
            "phanBo1": safe_str(ws.cell(row=row, column=16).value),
            "loaiThe1": safe_str(ws.cell(row=row, column=17).value),
            "loaiThe2": safe_str(ws.cell(row=row, column=18).value),
            "phanBo2": safe_str(ws.cell(row=row, column=19).value),
            "csNhap": safe_str(ws.cell(row=row, column=20).value),
            "tongSS": safe_num(ws.cell(row=row, column=21).value),
            "ptGym": safe_num(ws.cell(row=row, column=22).value),
            "pttl": safe_num(ws.cell(row=row, column=23).value),
            "daily": safe_num(ws.cell(row=row, column=24).value),
            "boi": safe_num(ws.cell(row=row, column=25).value),
            "ptTang": safe_num(ws.cell(row=row, column=26).value),
            "buoiPtCoDinh": safe_num(ws.cell(row=row, column=27).value),
            "soThang": safe_num(ws.cell(row=row, column=28).value),
            "tongSoThang": safe_num(ws.cell(row=row, column=29).value),
            "congThangBL": safe_num(ws.cell(row=row, column=30).value),
            "congThangTang": safe_num(ws.cell(row=row, column=31).value),
            "congNgayTang": safe_num(ws.cell(row=row, column=32).value),
            "goiTap": safe_str(ws.cell(row=row, column=33).value),
            "khungGioTap": safe_str(ws.cell(row=row, column=34).value),
            "chuyenNhuong": safe_str(ws.cell(row=row, column=35).value),
            "soHD": safe_str(ws.cell(row=row, column=36).value),
            "soThangBuoi": safe_str(ws.cell(row=row, column=37).value),
            "ngayBatDau": parse_date(ws.cell(row=row, column=38).value),
            "ngayKetThuc": parse_date(ws.cell(row=row, column=44).value),
            "tinhTrang": safe_str(ws.cell(row=row, column=45).value),
            "tongSoNgay": safe_num(ws.cell(row=row, column=46).value),
            "tongNgayDaDi": safe_num(ws.cell(row=row, column=47).value),
            "tongSoNgayConLai": safe_num(ws.cell(row=row, column=49).value),
            "quaTang": safe_str(ws.cell(row=row, column=50).value),
            "donGiaNiemYet": safe_num(ws.cell(row=row, column=51).value),
            "tongTienBaoLuu": safe_num(ws.cell(row=row, column=52).value),
            "tongTienBuoiDay": safe_num(ws.cell(row=row, column=53).value),
            "tongTien": safe_num(ws.cell(row=row, column=56).value),
            "mucUuDai": safe_num(ws.cell(row=row, column=57).value),
            "soTienKHGiam": safe_num(ws.cell(row=row, column=58).value),
            "tong": safe_num(ws.cell(row=row, column=59).value),
            "soTienDVPhaiThu": safe_num(ws.cell(row=row, column=60).value),
            "bnDatCoc": safe_str(ws.cell(row=row, column=61).value),
            "tenGoiNangCap": safe_str(ws.cell(row=row, column=62).value),
            "soTienKHNangCap": safe_str(ws.cell(row=row, column=63).value),
            "khDaDatCoc": safe_str(ws.cell(row=row, column=64).value),
            "tongDatCocNangCap": safe_str(ws.cell(row=row, column=65).value),
            "conPhaiThuKH": safe_str(ws.cell(row=row, column=66).value),
            "phanTramSaleDT": safe_num(ws.cell(row=row, column=67).value),
            "datCoc": safe_str(ws.cell(row=row, column=68).value),
            "ttPhanThieu": safe_str(ws.cell(row=row, column=69).value),
            "dongFull": safe_str(ws.cell(row=row, column=70).value),
            "phiPos": safe_str(ws.cell(row=row, column=71).value),
            "doanhThu": safe_str(ws.cell(row=row, column=72).value),
            "hinhThucTT": safe_str(ws.cell(row=row, column=73).value),
            "daNhanTM": safe_str(ws.cell(row=row, column=74).value),
            "ckChung": safe_str(ws.cell(row=row, column=75).value),
            "tongTienCKChung": safe_str(ws.cell(row=row, column=76).value),
            "trangThai": safe_str(ws.cell(row=row, column=77).value),
            "tongGoiTap": safe_str(ws.cell(row=row, column=78).value),
            "loai": safe_str(ws.cell(row=row, column=79).value),
            "luot": safe_str(ws.cell(row=row, column=80).value),
            "nguon": safe_str(ws.cell(row=row, column=81).value),
            "daNhapTimeSoft": safe_str(ws.cell(row=row, column=83).value),
            "ghiChu": safe_str(ws.cell(row=row, column=84).value),
        })
    return records
